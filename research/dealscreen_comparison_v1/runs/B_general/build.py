"""Workflow B builder: calls.csv, valuation, workbook. AI-SIMULATED / INTERNAL DEVELOPMENT."""
import csv, yaml, os
from openpyxl import Workbook
from openpyxl.styles import Font
H=os.path.dirname(os.path.abspath(__file__)); I=os.path.join(H,'..','..','inputs')
land={r['Program ID']:r for r in csv.DictReader(open(os.path.join(I,'landscape_cldn18.csv')))}
TK='https://www.takeda.com/newsroom/newsreleases/2025/innovent/'
AZ='https://www.astrazeneca.com/media-centre/press-releases/2023/astrazeneca-enters-license-agreement-with-kym-biosciences-for-cmg901-a-claudin-182-antibody-drug-conjugate.html'
AST='https://www.astellas.com/en/science/research-and-development/pipeline'
# id: call, reason, group, holder, avail, mkt, dev, rank, evidence
C={
'CL-01':('OUT','naked mAb; marketed (Vyloy)','zolbetuximab','Astellas (acquired Ganymed 2016)','HELD_BY_OTHERS','Marketed US (FDA approved 2024-10-18)','Active','','https://api.fda.gov/drug/drugsfda.json?search=openfda.generic_name:zolbetuximab "submission_status": "AP", "submission_status_date": "20241018"'),
'CL-02':('OUT','naked mAb; same molecule as zolbetuximab','zolbetuximab','Astellas','HELD_BY_OTHERS','Marketed US (as Vyloy)','Active','','https://api-evsrest.nci.nih.gov/api/v1/concept/ncit/C85475 "name":"IMAB362" (zolbetuximab concept)'),
'CL-03':('IN','CLDN18.2 MMAE ADC, Ph3 positive OS','sonesitatug_vedotin','AstraZeneca (global exclusive license from KYM)','HELD_BY_OTHERS','Not marketed','Active Ph3',' ',AZ+' "global exclusive licence agreement for CMG901"'),
'CL-04':('IN','same molecule as CL-03','sonesitatug_vedotin','AstraZeneca','HELD_BY_OTHERS','Not marketed','Active','',AZ+' "global exclusive licence agreement for CMG901"'),
'CL-05':('IN','same molecule as CL-03','sonesitatug_vedotin','AstraZeneca','HELD_BY_OTHERS','Not marketed','Active Ph3','',AZ+' "global exclusive licence agreement for CMG901"'),
'CL-06':('IN','landscape modality "Imaging agent" is WRONG: exatecan ADC (NCIt); Ph3 met primary endpoint','arcotatug_tavatecan','Takeda (ex-Greater China license 2025-10-21)','HELD_BY_OTHERS','Not marketed (NMPA review)','Active Ph3','',TK+' "license ... IBI363 and IBI343, worldwide outside of Greater China"'),
'CL-07':('IN','CLDN18.2 MMAE ADC Ph3 (China)','LM-302','LaNova (acquired by Sino Biopharm 2025); BMS ex-China license terminated/returned','NOT_EVIDENCED','Not marketed','Active Ph3 (China)','2','https://pharmaphorum.com/news/msd-az-partner-lanova-bought-sino-biopharm-951m "handed back rights last year"'),
'CL-08':('IN','CLDN18.2 exatecan-derivative ADC Ph3','SHR-A1904','Hengrui; Merck KGaA holds ex-mainland-China option (exercise status unknown)','HELD_BY_OTHERS','Not marketed','Active Ph3','','https://www.biospace.com/hengrui-pharma-partners-with-merck-kgaa-darmstadt-germany-to-advance-innovative-cancer-therapies "an exclusive option to develop ... SHR-A1904 worldwide, outside of mainland China"'),
'CL-09':('OUT','naked (ADCC-enhanced) mAb osemitamab','osemitamab','Transcenta','UNKNOWN','Not marketed','Active Ph2/3','','https://api-evsrest.nci.nih.gov/api/v1/concept/ncit/C173726 "humanized ... monoclonal antibody"'),
'CL-10':('IN','CLDN18.2xCD3 TCE Ph3','ASP2138','Astellas (in-house)','HELD_BY_OTHERS','Not marketed','Active Ph3','','placera.se Astellas PR 2026-08-25 "first patient has been dosed in a Phase 3 study evaluating ASP2138"'),
'CL-11':('IN','CLDN18.2 TOP1i ADC Ph1/2','XNW27011','Astellas (worldwide ex-Greater China, from Evopoint 2025)','HELD_BY_OTHERS','Not marketed','Active','',AST+' "ASP546C/XNW27011: Astellas has worldwide rights excluding China\'s mainland..."'),
'CL-12':('IN','same molecule as CL-11','XNW27011','Astellas','HELD_BY_OTHERS','Not marketed','Active','',AST+' "ASP546C/XNW27011"'),
'CL-13':('IN','CLDN18.2x4-1BB bispecific Ph2 1L gastric','givastomig','I-Mab + ABL Bio share ex-Greater China/Korea rights equally','NOT_EVIDENCED','Not marketed','Active Ph2','6','https://natlawreview.com/press-releases/i-mab-announces-publication-givastomig-monotherapy-data-clinical-cance "shares worldwide rights, excluding Greater China and South Korea, equally with ABL Bio"'),
'CL-14':('IN','CLDN18.2 ADC (science fits) but development discontinued','EO-3021','Elevation Oncology (ex-Greater China licence from CSPC); post-discontinuation rights unknown','UNKNOWN','Not marketed','Discontinued (2025-03)','','https://www.sec.gov/Archives/edgar/data/1783032/000110465925025878/tm259617d1_ex99-1.htm "we have decided to discontinue further development of EO-3021"'),
'CL-15':('IN','same molecule as EO-3021 (CSPC code)','EO-3021','CSPC (China); ex-China licensee Elevation','UNKNOWN','Not marketed','No active trial','','https://www.sec.gov/Archives/edgar/data/1783032/000155837022011277/elev-20220727xex99d1.htm "exclusive license agreement with CSPC Megalith"'),
'CL-16':('IN','CLDN18.2 ADC Ph3 G/GEJ; FDA ODD','ATG-022','Antengene (global rights, per 2024 interim results)','NOT_EVIDENCED','Not marketed','Active Ph3','1','https://www.1stoncology.com/blog/2024/page/503 (Antengene 2024 interim: ATG-022 listed among "global rights assets") ; NCT07795346'),
'CL-17':('IN','CLDN18.2xCD3 TCE Ph1/2','AZD5863','AstraZeneca','HELD_BY_OTHERS','Not marketed','Active','','landscape cell (Company: AstraZeneca); NCT06005493 "T cell-engaging bi-specific antibody that targets CLDN18.2"'),
'CL-18':('OUT','CAR-T (cell therapy excluded)','satri-cel','CARsgen','UNKNOWN','Marketed China (conditional NMPA 2026-06-22); not US','Active','','https://www.clinicaltrialvanguard.com/news/carsgens-satricabtagene-autoleucel-approved-for-gastric-cancer-in-c "conditionally approved"'),
'CL-19':('OUT','naked mAb','','CARsgen','UNKNOWN','Not marketed','No active trial','','landscape cell (Modality: Monoclonal antibody)'),
'CL-20':('IN','CLDN18.2xCD3 TCE Ph3 (China)','QLS31905','Qilu (no ex-China licensee found)','NOT_EVIDENCED','Not marketed','Active Ph3','7','https://www.biospace.com/qilu-pharmaceutical-announces-results-from-phase-i-clinical-study-for-its-novel-bispe "bispecific T cell engager (BiTE) that targets Claudin18.2"'),
'CL-21':('IN','likely QLS31905 under Anova code (same 31905 number, Qilu-family sponsor) - identity NOT confirmed','QLS31905','Anova Innovation (relationship to Qilu unverified)','UNKNOWN','Not marketed','Active Ph1b/2','','landscape cell; NCT07444541'),
'CL-22':('OUT','naked mAb','','FutureGen','UNKNOWN','Not marketed','Active','','landscape cell (Modality: Monoclonal antibody)'),
'CL-23':('OUT','target PD-L1 (noise)','','FutureGen','UNKNOWN','Not marketed','Active','','landscape cell (Target: PD-L1)'),
'CL-24':('IN','CLDN18.2 MMAE ADC Ph3 2L+ G/GEJ','JS107','Junshi / Junyan Hongshi (no ex-China licence found)','NOT_EVIDENCED','Not marketed','Active Ph3','3','NCT07284134 (Phase 3, recruiting); https://www.cancer.gov/publications/dictionaries/cancer-drug/def/anti-claudin-182-antibody-drug-conjugate-js107'),
'CL-25':('IN','CLDN18.2 ADC Ph1/2','RC118','RemeGen ("self-developed")','NOT_EVIDENCED','Not marketed','Registry status mixed','8','https://www.biospace.com/remegen-s-rc118-for-injection-targeting-claudin-18-2-in-patients-with-gastric-and-pan "Self-developed, RC118 is RemeGen\'s latest ... ADC"'),
'CL-26':('OUT','target PD-1-axis bispecific, not CLDN18.2 (noise); AbbVie ex-China','','AbbVie (ex-Greater China)','HELD_BY_OTHERS','Not marketed','Unknown','','https://news.abbvie.com/2026-01-12-AbbVie-and-RemeGen-Announce-Exclusive-Licensing-Agreement-to-Develop-A-Nove "exclusive rights ... RC148 outside of the Greater China"'),
'CL-27':('IN','CLDN18.2 ADC Ph1/2','SKB315','Kelun-Biotech (Merck returned global rights Aug 2024)','NOT_EVIDENCED','Not marketed','Active','5','https://synapse.patsnap.com/article/merck-co-invests-in-kelun-bispecific-adc-exits-cldn182-program (Merck gave back SKB315 rights; Kelun to "explore international markets" - paraphrase via search)'),
'CL-28':('OUT','naked CLDN18.2 mAb (PD-1 is combination partner)','','AskGene / Aosaikang','UNKNOWN','Not marketed','Active Ph3','','landscape cell (Modality: Monoclonal antibody)'),
'CL-29':('IN','CLDN18.2 TOP1i ADC Ph3 G/GEJ','BL-M05D1','Sichuan Baili / SystImmune (US affiliate); no outside licensee found','NOT_EVIDENCED','Not marketed','Active Ph3','4','NCT07518147 "BL-M05D1 With the Investigator\'s Choice ... CLDN18.2-Posit[ive]"'),
'CL-30':('OUT','naked blocking mAb; terminated','','Leap Therapeutics (acquired Flame)','UNKNOWN','Not marketed','Terminated','','https://www.sec.gov/Archives/edgar/data/1509745/000110465923004066/tm233557d1_ex99-1.htm "fully human monoclonal antibody"'),
'CL-31':('IN','CLDN18.2xCD47 bispecific Ph2','spevatamig','Phanes Therapeutics (MSD listed as collaborator - likely drug supply, unverified)','NOT_EVIDENCED','Not marketed','Active Ph2','','https://www.biospace.com/press-releases/phanes-therapeutics-announces-updated-phase-2-clinical-trial-results-o "bispecific antibody (bsAb) targeting claudin 18.2 and CD47"'),
'CL-32':('UNRESOLVED','ISAC (TLR7/8 agonist conjugate) - neither cytotoxic ADC nor bispecific; landscape says mAb','','Bolt Biotherapeutics','UNKNOWN','Not marketed','Active Ph1/2','','NCT06921837 "Immune stimulating antibody conjugate (ISAC)"'),
'CL-33':('IN','CLDN18.2 ADC; all trials terminated','','SOTIO','UNKNOWN','Not marketed','Terminated','','landscape cell (Status: terminated)'),
'CL-34':('OUT','naked mAb','','Zai Lab','UNKNOWN','Not marketed','Completed','','landscape cell'),
'CL-35':('OUT','mRNA-encoded mAb; discontinued','','BioNTech','UNKNOWN','Not marketed','Discontinued','','https://www.sec.gov/Archives/edgar/data/1776985/000177698523000050/bnt_2023q2ex99x1quarterlyr.htm "decided to discontinue the study"'),
'CL-36':('IN','CLDN18.2xCD3 BiTE (fits) but terminated','','Amgen','HELD_BY_OTHERS','Not marketed','Terminated','','landscape cell'),
'CL-37':('IN','CLDN18.2 MMAE ADC Ph1 (elatatug vedotin)','','TORL Biotherapeutics (US private)','NOT_EVIDENCED','Not marketed','Active Ph1','9','https://www.cismef.org/page/detail/en/NCI_CO_C185565 "Anti-CLDN18.2/MMAE ADC TORL-2-307-ADC"'),
'CL-38':('IN','landscape says mAb; actually ADC (ConjuAll, TOP1i) Ph1/2','','LigaChem','NOT_EVIDENCED','Not marketed','Active Ph1/2','','NCT07460375 "CLDN18.2-directed human monoclonal antibody (Ab) linked to a topoisomerase I inhibiting payload"'),
'CL-39':('IN','CLDN18.2xPD-L1 bispecific Ph1; registry UNKNOWN','','QureBio','UNKNOWN','Not marketed','Unknown','','landscape cell'),
'CL-40':('OUT','naked mAb','','Mabworks','UNKNOWN','Not marketed','Unknown','','landscape cell'),
'CL-41':('UNRESOLVED','CLDN18.2 antibody-IL-10 fusion (bifunctional), not clearly ADC/bispecific','','L&L Bio (reportedly seeking US partner - unverified)','UNKNOWN','Not marketed','Active Ph1/2','','WebSearch summary of L&L Bio materials; not verified by fetch'),
'CL-42':('OUT','CAR-T','','Legend Biotech','UNKNOWN','Not marketed','Active','','landscape cell'),
'CL-43':('IN','CLDN18.2 ADC Ph1; registry UNKNOWN','','Chia Tai Tianqing','UNKNOWN','Not marketed','Unknown','','NCT05867563 "antibody-drug conjugate targeting CLDN18.2"'),
'CL-44':('IN','CLDN18.2xCD47 bispecific Ph1','','Sumgen','NOT_EVIDENCED','Not marketed','Active Ph1','','https://www.cancer.gov/publications/dictionaries/cancer-drug/def/anti-cldn182-cd47-bispecific-antibody-sg1906'),
'CL-45':('UNRESOLVED','format not determined; terminated','','Doer Biologics','UNKNOWN','Not marketed','Terminated','','landscape cell (Modality blank)'),
'CL-46':('IN','CLDN18.2xCDH17 bispecific ADC Ph1/2','','VelaVigo','NOT_EVIDENCED','Not marketed','Active Ph1/2','10','NCT07700160; WebSearch summary "CLDN18.2 and CDH17-directed Bi-specific Antibody Drug Conjugate"'),
'CL-47':('OUT','CAR-T','','Innovent','UNKNOWN','Not marketed','Completed','','landscape cell'),
'CL-48':('OUT','TAC T-cell therapy (cell therapy)','','Triumvira','UNKNOWN','Not marketed','Active','','NCT05862324 "TAC T-cells"'),
'CL-49':('OUT','CAR-T','','Immunofoco','UNKNOWN','Not marketed','Unknown','','biospace "autologous CLDN18.2-targeting CAR-T"'),
'CL-50':('OUT','naked mAb; terminated','','LaNova','UNKNOWN','Not marketed','Terminated','','landscape cell'),
'CL-51':('IN','CLDN18.2xCD47 bispecific Ph1; registry UNKNOWN','','Akeso','UNKNOWN','Not marketed','Unknown','','landscape cell'),
'CL-52':('OUT','radioligand - not an in-scope modality','','SmartNuclide','UNKNOWN','Not marketed','Active','','landscape cell'),
'CL-53':('UNRESOLVED','target/modality unknown (curator not_found)','','Yiling','UNKNOWN','Not marketed','Unknown','','landscape cell (Target blank)'),
'CL-54':('OUT','naked mAb; withdrawn','','SparX','UNKNOWN','Not marketed','Withdrawn','','landscape cell'),
'CL-55':('IN','CLDN18.2 ADC Ph1/2a (Korea)','','Dong-A ST / AbTis','NOT_EVIDENCED','Not marketed','Active Ph1','','NCT07481357; WebSearch summary (dailypharm) "ADC targeting Claudin18.2"'),
'CL-56':('OUT','CAR-NK','','Beijing Biotech','UNKNOWN','Not marketed','Active','','landscape cell'),
'CL-57':('OUT','PD-1/TIGIT, not CLDN18.2 (noise)','','AstraZeneca','HELD_BY_OTHERS','Not marketed','Active','','landscape cell'),
'CL-58':('OUT','HER2 bispecific, not CLDN18.2 (noise); sponsor field wrong (AACR)','','Jazz/BeOne','HELD_BY_OTHERS','Marketed US (other indication)','Active','','landscape cell'),
'CL-59':('OUT','PD-1/CTLA-4 (noise)','','AstraZeneca','HELD_BY_OTHERS','Not marketed','Active','','landscape cell'),
'CL-60':('OUT','PD-1/CTLA-4 (noise)','','Qilu','UNKNOWN','Not marketed','Active','','landscape cell'),
'CL-61':('OUT','imaging agent','','Wuxi No.4 Hospital','UNKNOWN','Not marketed','Unknown','','landscape cell'),
'CL-62':('OUT','imaging agent','','RenJi Hospital','UNKNOWN','Not marketed','Active','','landscape cell'),
}
assert len(C)==62 and set(C)==set(land)
cols=['program_id','scientific_call','reason','identity_group','us_rights_holder','availability','marketing_status','development_status','shortlist_rank','evidence']
rows=[dict(zip(cols,[k,*[str(x).strip() for x in v]])) for k,v in C.items()]
with open(os.path.join(H,'calls.csv'),'w',newline='') as f:
    w=csv.DictWriter(f,cols); w.writeheader(); w.writerows(rows)

# ---- valuation (HYPOTHETICAL) ----
M=yaml.safe_load(open(os.path.join(I,'valuation_market_inputs_hypothetical.yaml')))
mk=M['market_us_cldn18_gastric_gej']; dv=M['development_hypothetical_by_next_phase']; dt=M['deal_terms_hypothetical']
EBIT=0.40  # ANALYST ASSUMPTION: operating margin on net sales (pre-royalty)
SC={'bear':(0.10,0.8),'base':(mk['peak_penetration'],1.0),'bull':(0.20,1.2)}  # ANALYST: penetration, PoS multiplier
def rnpv(phases,scen,deal):
    pen,pm=SC[scen]; r=mk['discount_rate']; t=0; p=1.0; v=0.0
    for ph in phases:
        d=dv[ph]; v-=p*d['cost_millions']/(1+r)**t  # cost paid at phase start, risked by prior success
        p*=min(1.0,d['success_probability']*pm); t+=d['duration_years']
    peak=mk['addressable_patients_annual']*mk['net_price_per_patient_usd']*pen/1e6
    roy=dt['license']['royalty_rate'] if deal=='license' else 0.0
    for y in range(1,mk['patent_life_years']+1):
        s=peak*min(1.0,y/mk['years_to_peak'])
        v+=p*s*(EBIT-roy)/(1+r)**(t+y-0.5)
    if deal=='license': v+=-dt['license']['upfront_cost_millions']-p*dt['license']['milestones_total_millions']/(1+r)**t
    else: v-=dt['acquisition']['upfront_cost_millions']
    return round(v,1),round(p,4),round(peak,1)
VAL={'CL-16 ATG-022':['phase_3','nda_bla'],'CL-07 LM-302':['phase_3','nda_bla'],'CL-27 SKB315':['phase_2','phase_3','nda_bla']}
vr=[]
for n,ph in VAL.items():
    for deal in ('license','acquisition'):
        for s in SC: v,p,pk=rnpv(ph,s,deal); vr.append([n,'+'.join(ph),deal,s,pk,p,v])
with open(os.path.join(H,'valuation_table.csv'),'w',newline='') as f:
    w=csv.writer(f); w.writerow(['program','remaining_phases','deal','scenario','peak_sales_musd','cum_pos','buyer_npv_musd']); w.writerows(vr)

wb=Workbook(); ws=wb.active; ws.title='README'
for l in ['CLDN18.2 G/GEJ screen - Workflow B (general Claude). AI-SIMULATED / INTERNAL DEVELOPMENT.',
 'Human review: NOT COMPLETED BY HUMAN. Valuation inputs HYPOTHETICAL.','Sheets: Calls (62 rows), Shortlist, Valuation, Identity']: ws.append([l])
ws=wb.create_sheet('Calls'); ws.append(cols+['program','company','modality(landscape)','phase(landscape)'])
for r in rows:
    L=land[r['program_id']]; ws.append([r[c] for c in cols]+[L['Program'],L['Company'],L['Modality'],L['Highest Phase']])
ws=wb.create_sheet('Shortlist'); ws.append(['rank','program_id','program','holder','availability','key risk / open question'])
RISK={'CL-16':'Confirm no ex-China deal since mid-2025; Ph3 data timing','CL-07':'Sino Biopharm ownership; is it seeking a partner after BMS return?','CL-24':'Junshi partnering intent unknown; MMAE crowded','CL-29':'Baili may keep for SystImmune/own US dev','CL-27':'Ph1/2 only; Merck returned (why?)','CL-13':'Two rights holders (I-Mab+ABL); 4-1BB tox','CL-20':'TCE not ADC; ANO31905 identity','CL-25':'Registry status UNKNOWN; partnering unknown','CL-37':'Ph1; private co','CL-46':'Ph1; bispecific ADC novelty'}
for r in sorted([r for r in rows if r['shortlist_rank']],key=lambda r:int(r['shortlist_rank'])):
    ws.append([int(r['shortlist_rank']),r['program_id'],land[r['program_id']]['Program'],r['us_rights_holder'],r['availability'],RISK[r['program_id']]])
ws=wb.create_sheet('Valuation'); ws.append(['program','remaining_phases','deal','scenario','peak_sales_musd','cum_pos','buyer_npv_musd']); [ws.append(x) for x in vr]
ws.append([]); ws.append(['ASSUMPTIONS: HYPOTHETICAL shared inputs + analyst EBIT margin 40%, bear pen 10%/PoS x0.8, bull pen 20%/PoS x1.2; see valuation.md'])
ws=wb.create_sheet('Identity'); ws.append(['identity_group','program_ids'])
g={}
for r in rows:
    if r['identity_group']: g.setdefault(r['identity_group'],[]).append(r['program_id'])
for k,v in g.items(): ws.append([k,'; '.join(v)])
for s in wb: 
    for c in s[1]: c.font=Font(bold=True)
wb.save(os.path.join(H,'cldn18_screen_workflowB.xlsx'))
print(sum(r['scientific_call']=='IN' for r in rows),'IN'); [print(x) for x in vr]
