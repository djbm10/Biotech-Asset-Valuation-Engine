"""Full workflow: landscape → screen → research → review → valuation (BVE CLI) → XLSX + memo."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import pytest

from bve.cli.dealscreen import run_demo
from bve.dealscreen.valuation_handoff import (
    build_config,
    check_readiness,
    read_summary,
    run_valuation,
)
from bve.dealscreen.workbook import WorkbookExists, export_workbook
from bve.dealscreen.workspace import Workspace

N_SIMS = 300

REQUIRED_SHEETS = ["Landscape", "Screening", "Diligence", "Finalists", "Valuation Inputs",
                   "Valuation Scenarios", "Deal Comparison", "Audit Trail"]


@pytest.fixture(scope="module")
def demo(tmp_path_factory, demo_dir):
    from bve.se.ontology.targets import reset_resolver_cache

    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("BVE_SE_ONTOLOGY_SNAPSHOT", str(tmp_path_factory.mktemp("none") / "absent"))
        reset_resolver_cache()
        root = tmp_path_factory.mktemp("e2e") / "ws"
        args = argparse.Namespace(workspace=str(root), examples=str(demo_dir),
                                  landscape="landscape_200.csv",
                                  target_map=str(demo_dir / "target_map.yaml"), n_sims=N_SIMS)
        report = run_demo(args)
        yield Workspace(root), report
    reset_resolver_cache()


def test_end_to_end_report(demo):
    _, report = demo
    m = report["manifest"]
    assert m["rows_read"] == 200 == m["rows_imported"] + m["rows_rejected"] + m["rows_exact_duplicate"]
    for key in ("funnel_initial", "funnel_after_review"):
        f = report[key]
        assert f["eligible"] + f["unresolved"] + f["excluded"] == f["assets_screened"] == m["assets"]
    assert report["determinism"] == {"import_json_identical": True,
                                     "initial_screen_byte_identical_on_rerun": True}
    after = report["evaluation_after_review"]
    assert after["false_exclusions"] == [] and after["false_eligible"] == []
    assert after["identity_merge_errors"] == [] and after["rights_available_without_explicit_statement"] == []
    refused = report["fixture_phase_2"]["refused"]
    assert [r["subject"] for r in refused] and all("EXCLUDED" in r["reason"] for r in refused)


def test_finalists_are_only_logged_approvals(demo):
    ws, _ = demo
    log = ws.decision_log()
    assert 0 < len(log.state.finalists) <= 10
    assert all(rec["reviewer_kind"] == "SYNTHETIC_FIXTURE" for rec in log.state.finalists.values())


def test_missing_assumptions_mean_not_ready_and_no_valuation(demo):
    ws, _ = demo
    entries = ws.valuations()["assets"]
    not_ready = [e for e in entries.values() if not e["readiness"]["ready"]]
    assert not_ready and all(e["runs"] == [] for e in not_ready)
    partial = [e for e in not_ready if 0 < len(e["readiness"]["missing"]) < 11]
    assert partial and any("net_price" in m for m in partial[0]["readiness"]["missing"])
    assert len([e for e in entries.values() if e["runs"]]) == 3


def test_workbook_values_equal_saved_valuation_outputs(demo):
    from openpyxl import load_workbook

    _, report = demo
    wb = load_workbook(report["deliverables"]["workbook"], read_only=True)
    assert all(name in wb.sheetnames for name in REQUIRED_SHEETS)
    for name in REQUIRED_SHEETS:
        assert wb[name].max_row > 2, f"{name} is empty"
    rows = list(wb["Valuation Scenarios"].iter_rows(min_row=3, values_only=True))
    assert rows
    for row in rows:
        summary = read_summary(row[13])
        assert row[3:7] == (summary["bear_rnpv_millions"], summary["base_rnpv_millions"],
                            summary["bull_rnpv_millions"], summary["rnpv_millions"])
    deals = list(wb["Deal Comparison"].iter_rows(min_row=3, values_only=True))
    assert {r[2] for r in deals} == {"acquisition", "license"}
    landscape = list(wb["Landscape"].iter_rows(min_row=3, values_only=True))
    assert len(landscape) == 200
    statuses = [r[0] for r in landscape]
    assert sum(s.startswith("REJECTED") for s in statuses) == 2
    assert sum(s.startswith("EXACT DUPLICATE") for s in statuses) == 3
    wb.close()


def test_export_never_overwrites_an_analyst_workbook(demo, tmp_path):
    from openpyxl import Workbook

    ws, report = demo
    analyst = tmp_path / "analyst.xlsx"
    book = Workbook()
    book.active["A1"] = "=SUM(1,2)"
    book.save(analyst)
    before = analyst.read_bytes()
    with pytest.raises(WorkbookExists, match="not a DealScreen workbook"):
        export_workbook(ws, analyst)
    assert analyst.read_bytes() == before
    with pytest.raises(WorkbookExists, match="replace_own"):
        export_workbook(ws, report["deliverables"]["workbook"])


def test_memo_is_bounded_and_labelled(demo):
    _, report = demo
    text = Path(report["deliverables"]["markdown"]).read_text()
    assert "UNVALIDATED" in text and "SYNTHETIC reviewer" in text and "NOT READY FOR VALUATION" in text
    assert len(text.split()) < 600
    assert "we recommend" not in text.casefold()


def test_prototype_path_reconciles_with_bve_asset(demo, tmp_path):
    """With no deal terms, the prototype CLI and bve-asset give the same engine output."""
    import yaml

    ws, _ = demo
    asset_id = next(a for a, e in ws.valuations()["assets"].items() if e["runs"])
    asset = next(a for a in ws.assets() if a.asset_id == asset_id)
    assumptions = yaml.safe_load((ws.inputs / "valuation_assumptions.yaml").read_text())
    block = assumptions["assets"][asset.external_id]
    assert check_readiness(asset, block).ready
    config = build_config(asset, block, indication_label="Ulcerative colitis", deal=None)
    via_bve_asset = run_valuation(config, tmp_path / "a", case="standalone", n_sims=N_SIMS, seed=7,
                                  asset_id=asset_id)
    config["deal_economics"] = {"upfront_cost_millions": 0.0}  # forces the prototype path, no terms
    via_prototype = run_valuation(config, tmp_path / "b", case="standalone", n_sims=N_SIMS, seed=7,
                                  asset_id=asset_id)
    assert via_prototype.prototype and not via_bve_asset.prototype
    keys = ("rnpv_millions", "bull_rnpv_millions", "base_rnpv_millions", "bear_rnpv_millions",
            "mc_mean_millions", "cumulative_success_probability")
    assert {k: via_bve_asset.summary[k] for k in keys} == {k: via_prototype.summary[k] for k in keys}
    saved = json.loads(Path(via_bve_asset.output_json).read_text())
    assert saved["outputs"]["rnpv"]["rnpv_millions"] == via_bve_asset.summary["rnpv_millions"]
