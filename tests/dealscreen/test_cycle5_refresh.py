"""Cycle 5: analyst decisions survive a vendor refresh; changes and orphans are reported."""

from __future__ import annotations

import csv
import random
from pathlib import Path

import pytest

from bve.dealscreen.review import Action
from bve.dealscreen.workspace import Workspace

REAL = Path(__file__).resolve().parents[2] / "examples" / "dealscreen_real"
TARGETS = {t.casefold(): t for t in ["BTK", "C5", "CD19", "CD38", "CD40", "CFB", "FCGRT", "IL6R",
                                      "MS4A1", "TNFSF13B"]} | {"bcma": "TNFRSF17"}
PROFILE, BUYER = REAL / "profile_public_landscape.yaml", REAL / "buyer_fcrn_gmg.yaml"


def norm(text: str) -> str | None:
    return TARGETS.get(text.strip().casefold())


def _rows() -> list[dict[str, str]]:
    return list(csv.DictReader((REAL / "landscape_fcrn_gmg.csv").open()))


def _write(path: Path, rows: list[dict[str, str]]) -> Path:
    with path.open("w", newline="") as h:
        w = csv.DictWriter(h, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return path


@pytest.fixture
def ws(tmp_path, monkeypatch):
    from bve.se.ontology.targets import reset_resolver_cache

    monkeypatch.setenv("BVE_SE_ONTOLOGY_SNAPSHOT", str(tmp_path / "absent"))
    reset_resolver_cache()
    w = Workspace(tmp_path / "ws")
    w.import_landscape(REAL / "landscape_fcrn_gmg.csv", PROFILE, BUYER, target_normalizer=norm)
    w.screen(None)
    yield w
    reset_resolver_cache()


def _disp(ws, ext):
    ids = {a.external_id: a.asset_id for a in ws.assets()}
    return next(s["disposition"] for s in ws.screening()["screens"] if s["asset_id"] == ids[ext])


def _reject_target(ws, ext):
    asset = next(a for a in ws.assets() if a.external_id == ext)
    claim = asset.claims_for("target")[0]
    ws.decision_log().record(action=Action.REJECT_CLAIM, subject_id=claim.claim_id, new_state="REJECTED",
                             reviewer_id="analyst-1", rationale="wrong target", context=ws.decision_context(norm))
    return claim.claim_id


def test_f1_f3_reordered_reexport_keeps_every_decision(ws, tmp_path):
    _reject_target(ws, "RW-03")
    ws.screen(None)
    assert _disp(ws, "RW-03") == "UNRESOLVED"
    before = {s["asset_id"]: s["disposition"] for s in ws.screening()["screens"]}
    rows = _rows()
    random.Random(7).shuffle(rows)
    ws.import_landscape(_write(tmp_path / "shuffled.csv", rows), PROFILE, BUYER, target_normalizer=norm)
    ws.screen(None)
    assert _disp(ws, "RW-03") == "UNRESOLVED"  # the reproduced regression: was ELIGIBLE before the fix
    assert {s["asset_id"]: s["disposition"] for s in ws.screening()["screens"]} == before
    report = ws.refresh()
    assert report["orphaned_decisions"] == [] and report["field_changes"] == {}


def test_f2_f4_changed_value_orphans_the_decision_and_is_reported(ws, tmp_path):
    claim_id = _reject_target(ws, "RW-03")
    rows = _rows()
    for row in rows:
        if row["Program ID"] == "RW-03":
            row["Target"] = "FCGRT; CD19"  # vendor changed the value
    rows = [r for r in rows if r["Program ID"] != "RW-40"]  # and dropped a program
    ws.import_landscape(_write(tmp_path / "changed.csv", rows), PROFILE, BUYER, target_normalizer=norm)
    report = ws.refresh()
    assert [o["subject"] for o in report["orphaned_decisions"]] == [claim_id]
    assert any(c["field"] == "target" for c in report["field_changes"]["RW-03 Rozanolixizumab"])
    assert report["assets_removed"] == ["RW-40 NMD670"]
    assert len(list((ws.root / "import_history").glob("import_*.json"))) == 1
    ws.screen(None)
    # The rejection of the OLD value is not transferred to the new value.
    assert _disp(ws, "RW-03") == "ELIGIBLE"  # new value FCGRT; CD19 passes ANY(FCGRT) on its own evidence
    assert claim_id in ws.decision_log().state.claims  # history kept


def test_f5_workbook_lists_orphaned_decisions(ws, tmp_path):
    from openpyxl import load_workbook

    from bve.dealscreen.workbook import export_workbook

    _reject_target(ws, "RW-03")
    rows = _rows()
    for row in rows:
        if row["Program ID"] == "RW-03":
            row["Target"] = "CD19"
    ws.import_landscape(_write(tmp_path / "c.csv", rows), PROFILE, BUYER, target_normalizer=norm)
    ws.screen(None)
    sheet = load_workbook(export_workbook(ws, tmp_path / "o.xlsx"), read_only=True)["Diligence"]
    assert any(r[2] == "ORPHANED_DECISION_AFTER_REFRESH" for r in sheet.iter_rows(values_only=True))


# -- regressions from the Cycles 3-5 adversarial review (AI reviewer) --------------------
def test_review_h1_uncached_asserted_target_disables_the_fast_path(ws, monkeypatch):
    from bve.dealscreen import workspace as wsmod

    ids = {a.external_id: a.asset_id for a in ws.assets()}
    ws.decision_log().record(
        action=Action.ASSERT_FACT, subject_id=ids["RW-08"], new_state="ASSERTED", reviewer_id="analyst-1",
        rationale="10-Q", detail={"field": "target", "value": "FcRn", "source_url": "https://www.sec.gov/x",
                                  "evidence_span": "anti-FcRn strategy based on its lead asset, IMVT-1402"},
        context=ws.decision_context(lambda t: "FCGRT" if t == "FcRn" else norm(t)))
    assert not ws.can_screen_without_ontology()  # 'FcRn' is not in the import-time cache
    monkeypatch.setattr(wsmod, "ontology_target_normalizer", lambda: (lambda t: "FCGRT" if t == "FcRn" else norm(t)))
    result = ws.screen(None)
    screen = next(s for s in result.screens if s.asset_id == ids["RW-08"])
    assert screen.target_symbols == ["FCGRT"]  # not silently "no target"


def test_review_l2_changed_ontology_version_disables_the_fast_path(ws):
    import json

    meta = ws.root / "target_cache_meta.json"
    meta.write_text(json.dumps({"ontology_version": "some-other-release"}))
    assert not ws.can_screen_without_ontology()


def test_review_m2_same_value_from_two_sources_keeps_decisions_on_the_right_source(tmp_path, monkeypatch):
    from bve.dealscreen.importer import import_landscape
    from bve.dealscreen.mapping import MappingProfile

    profile = MappingProfile.from_yaml(PROFILE)
    rows = [r for r in _rows() if r["Program ID"] == "RW-03"]
    a, b = dict(rows[0]), dict(rows[0])
    a["Sources"], b["Sources"] = "https://source-a.example", "https://source-b.example"
    first = import_landscape(_write(tmp_path / "1.csv", [a, b]), profile).assets[0]
    second = import_landscape(_write(tmp_path / "2.csv", [b, a]), profile).assets[0]
    by_source = lambda asset: {c.source_reference: c.claim_id for c in asset.claims_for("target")}
    assert by_source(first) == by_source(second)
    assert len(set(by_source(first).values())) == 2


def test_review_m3_legacy_ids_are_reported_as_legacy(ws, tmp_path):
    from bve.dealscreen.refresh import refresh_report
    from bve.dealscreen.review import ReviewState

    state = ReviewState(claims={"ds:abc:r5:target": "REJECTED"})
    report = refresh_report(ws.import_result(), ws.import_result(), state, [], [])
    assert "legacy row-based claim id" in report["orphaned_decisions"][0]["reason"]


def test_review_m4_decided_asset_with_changed_source_needs_review(ws, tmp_path):
    ids = {a.external_id: a.asset_id for a in ws.assets()}
    ws.decision_log().record(action=Action.REQUEST_DILIGENCE, subject_id=ids["RW-04"], new_state="REQUESTED",
                             reviewer_id="a", rationale="r", detail={"question": "rights?"},
                             context=ws.decision_context(norm))
    rows = _rows()
    for row in rows:
        if row["Program ID"] == "RW-04":
            row["Sources"] = row["Sources"] + "; https://new-source.example"
    ws.import_landscape(_write(tmp_path / "s.csv", rows), PROFILE, BUYER, target_normalizer=norm)
    assert ws.refresh()["decided_assets_with_changed_content"] == ["RW-04 Nipocalimab"]


def test_review_m5_failed_refresh_leaves_the_workspace_untouched(ws, tmp_path):
    from bve.dealscreen.review import LogTampered

    ids = {a.external_id: a.asset_id for a in ws.assets()}
    ws.decision_log().record(action=Action.REQUEST_DILIGENCE, subject_id=ids["RW-04"], new_state="REQUESTED",
                             reviewer_id="a", rationale="r", context=ws.decision_context(norm))
    ws.decisions_path.write_text(ws.decisions_path.read_text().replace('"reviewer_id": "a"', '"reviewer_id": "b"'))
    before = (ws.root / "import.json").read_bytes()
    rows = [r for r in _rows() if r["Program ID"] != "RW-40"]
    with pytest.raises(LogTampered):
        ws.import_landscape(_write(tmp_path / "x.csv", rows), PROFILE, BUYER, target_normalizer=norm)
    assert (ws.root / "import.json").read_bytes() == before
    assert not (ws.root / "import_history").exists()
